"""Bounded Office text readers and new-document exporters.

Optional source dependencies are bundled in the desktop application. These
operations do not execute macros, formulas, model code or shell commands.
"""
from __future__ import annotations
import csv
import io
import math
from pathlib import Path, PurePosixPath
import re
import tempfile
import uuid
import zipfile
from .paths import Workspace, PolicyError

FORMATS = ('md', 'docx', 'xlsx', 'pptx')
MAX_TEXT = 200000


def validate_office(path: Path):
    if path.stat().st_size > 16 * 1024 * 1024:
        raise PolicyError('Document exceeds 16 MiB')
    from defusedxml.ElementTree import fromstring
    with zipfile.ZipFile(path) as archive:
        parts = archive.infolist()
        if len(parts) > 5000 or sum(p.file_size for p in parts) > 32 * 1024 * 1024:
            raise PolicyError('Office archive exceeds expanded size limit')
        seen = set()
        for part in parts:
            name = part.filename
            if name in seen or '\\' in name or PurePosixPath(name).is_absolute() or '..' in PurePosixPath(name).parts:
                raise PolicyError('Unsafe or duplicate archive member')
            seen.add(name)
            if part.flag_bits & 1 or part.file_size > 8 * 1024 * 1024:
                raise PolicyError('Encrypted or oversized Office part')
            if 'vbaproject' in name.lower() or '/embeddings/' in name.lower():
                raise PolicyError('Macros and embedded objects are unsupported')
            if name.endswith(('.xml', '.rels')):
                try:
                    fromstring(archive.read(part), forbid_dtd=True, forbid_entities=True, forbid_external=True)
                except Exception as exc:
                    raise PolicyError('Invalid or unsafe Office XML') from exc
        if '[Content_Types].xml' not in seen:
            raise PolicyError('Not an Office Open XML document')


def table_markdown(rows):
    def cell(value):
        return str(value if value is not None else '').replace('|', '\\|').replace('\n', ' / ')
    if not rows:
        return ''
    width = max(len(row) for row in rows)
    lines = ['| ' + ' | '.join(cell(v) for v in list(row) + [''] * (width - len(row))) + ' |' for row in rows]
    lines.insert(1, '| ' + ' | '.join('---' for _ in range(width)) + ' |')
    return '\n'.join(lines)


def markdown_table(content):
    lines = content.splitlines()
    for i in range(len(lines) - 1):
        if '|' not in lines[i] or not re.fullmatch(r'[\s|:\-]+', lines[i + 1]) or '-' not in lines[i + 1]:
            continue
        rows = [[c.strip() for c in lines[i].strip().strip('|').split('|')]]
        for line in lines[i + 2:]:
            if '|' not in line or not line.strip():
                break
            rows.append([c.strip() for c in line.strip().strip('|').split('|')])
        return rows
    return None


class DocumentTools:
    def __init__(self, workspace: Workspace):
        self.ws = workspace

    @staticmethod
    def available():
        try:
            import docx, openpyxl, pptx, defusedxml
            return True
        except ImportError:
            return False

    def read(self, path: str, offset: int = 0, limit: int = 12000) -> dict:
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 16000:
            raise PolicyError('Invalid text window')
        file = self.ws.path(path)
        if not file.is_file():
            raise PolicyError('Document not found')
        ext = file.suffix.lower()
        sections = []
        note = ''
        if ext in ('.md', '.txt', '.csv', '.json', '.log'):
            content = self.ws.read(path)
        elif ext in ('.docx', '.xlsx', '.pptx'):
            if not self.available():
                raise PolicyError('Office libraries missing; use the desktop bundle or requirements-office.txt')
            validate_office(file)
            if ext == '.docx':
                from docx import Document
                from docx.table import Table
                for block in Document(file).iter_inner_content():
                    if isinstance(block, Table):
                        if len(block.rows) * len(block.columns) > 20000:
                            raise PolicyError('Word table too large')
                        sections.append(table_markdown([[c.text for c in r.cells] for r in block.rows]))
                    else:
                        name = block.style.name if block.style else ''
                        prefix = '#' * min(6, int(name[-1])) + ' ' if name.startswith('Heading ') and name[-1:].isdigit() else ''
                        sections.append(prefix + block.text)
                note = '正文与表格文本；不包含图片识别、批注、页眉页脚或修订还原。'
            elif ext == '.xlsx':
                from openpyxl import load_workbook
                book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
                try:
                    cells = 0
                    for sheet in book.worksheets:
                        sheet.reset_dimensions()
                        rows = []
                        for row in sheet.iter_rows():
                            cells += len(row)
                            if cells > 20000 or len(rows) >= 5000:
                                raise PolicyError('Workbook exceeds 20,000 cells / 5,000 rows per sheet')
                            rows.append([c.value for c in row])
                        sections.append('## ' + sheet.title + '\n\n' + table_markdown(rows))
                finally:
                    book.close()
                note = '显示值和公式文本，不执行或计算公式，不加载外部链接。'
            else:
                from pptx import Presentation
                presentation = Presentation(file)
                if len(presentation.slides) > 200:
                    raise PolicyError('Presentation exceeds 200 slides')
                def extract(shapes):
                    out = []
                    for shape in shapes:
                        if hasattr(shape, 'shapes'):
                            out.extend(extract(shape.shapes))
                        elif shape.has_text_frame:
                            out.append(shape.text)
                        elif shape.has_table:
                            out.append(table_markdown([[c.text for c in row.cells] for row in shape.table.rows]))
                    return out
                for index, slide in enumerate(presentation.slides, 1):
                    sections.append('## Slide ' + str(index) + '\n\n' + '\n\n'.join(extract(slide.shapes)))
                note = '按页提取文字和表格，不包含图片 OCR、动画或备注。'
            content = '\n\n'.join(sections)
        else:
            raise PolicyError('Unsupported format; use docx/xlsx/pptx/md/txt/csv/json/log without macros')
        if len(content) > MAX_TEXT:
            raise PolicyError('Extracted text exceeds 200,000 characters; split the document')
        end = min(len(content), offset + limit)
        return {'path': path, 'format': ext[1:], 'content': content[offset:end], 'total_chars': len(content),
                'offset': offset, 'next_offset': end if end < len(content) else None,
                'truncated': end < len(content), 'note': note}

    def create(self, format: str, content: str, title: str = 'Document') -> dict:
        if format not in FORMATS or not isinstance(content, str) or not 1 <= len(content) <= MAX_TEXT:
            raise PolicyError('Choose md/docx/xlsx/pptx and 1..200,000 characters')
        if not isinstance(title, str) or not 1 <= len(title) <= 100:
            raise PolicyError('Invalid title')
        slug = re.sub(r'[^\w -]', '', title).strip(' .')[:50] or 'Document'
        relative = f'exports/{uuid.uuid4().hex[:8]}-{slug}.{format}'
        out = self.ws.path(relative)
        out.parent.mkdir(parents=True, exist_ok=True)
        if format == 'md':
            return self.ws.write(relative, content) | {'format': 'md'}
        if not self.available():
            raise PolicyError('Office libraries are unavailable')
        rows = markdown_table(content)
        with tempfile.TemporaryDirectory(prefix='.agent-doc-', dir=out.parent) as tmp:
            target = Path(tmp) / ('document.' + format)
            if format == 'docx':
                from docx import Document
                from docx.shared import Inches, Pt
                doc = Document()
                section = doc.sections[0]
                section.top_margin = section.bottom_margin = Inches(.75)
                section.left_margin = section.right_margin = Inches(.8)
                normal = doc.styles['Normal']
                normal.font.name = 'Calibri'; normal.font.size = Pt(11)
                normal.paragraph_format.space_after = Pt(7)
                doc.add_heading(title, 0)
                lines = content.splitlines(); index = 0
                while index < len(lines):
                    line = lines[index]
                    if rows and index + 1 < len(lines) and '|' in line and re.fullmatch(r'[\s|:\-]+', lines[index + 1]):
                        width = max(len(row) for row in rows)
                        if width > 12 or len(rows) > 200:
                            raise PolicyError('Word table exceeds 12 columns / 200 rows')
                        table = doc.add_table(rows=0, cols=width); table.style = 'Light Shading Accent 1'
                        for values in rows:
                            cells = table.add_row().cells
                            for j, value in enumerate(values): cells[j].text = value
                        index += len(rows) + 1; rows = None
                        continue
                    heading = re.match(r'^(#{1,6})\s+(.+)', line)
                    if heading: doc.add_heading(heading[2], min(4, len(heading[1])))
                    elif line.startswith(('- ', '* ')): doc.add_paragraph(line[2:], style='List Bullet')
                    elif line.strip(): doc.add_paragraph(line)
                    index += 1
                doc.save(target)
            elif format == 'xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.utils import get_column_letter
                rows = rows or list(csv.reader(io.StringIO(content)))
                if len(rows) > 5000 or sum(map(len, rows)) > 20000 or max(map(len, rows), default=0) > 100:
                    raise PolicyError('Spreadsheet export exceeds size limits')
                book = Workbook(); sheet = book.active; sheet.title = 'Data'; sheet.freeze_panes = 'A2'
                for i, values in enumerate(rows, 1):
                    for j, value in enumerate(values, 1):
                        cell = sheet.cell(i, j)
                        if re.fullmatch(r'-?(?:0|[1-9]\d{0,13})(?:\.\d+)?', value or ''):
                            cell.value = float(value) if '.' in value else int(value)
                        else:
                            cell.value = value; cell.data_type = 's'
                        cell.alignment = Alignment(vertical='top', wrap_text=True)
                        if i == 1:
                            cell.fill = PatternFill('solid', fgColor='234E3A')
                            cell.font = Font(color='FFFFFF', bold=True, size=11)
                    sheet.row_dimensions[i].height = 30 if i == 1 else 24
                for j in range(1, sheet.max_column + 1):
                    sheet.column_dimensions[get_column_letter(j)].width = min(40, max(12, max(len(str(sheet.cell(i, j).value or '')) for i in range(1, sheet.max_row + 1)) + 2))
                sheet.auto_filter.ref = sheet.dimensions
                book.save(target); book.close()
            else:
                from pptx import Presentation
                from pptx.util import Inches, Pt
                from pptx.dml.color import RGBColor
                presentation = Presentation(); presentation.slide_width = Inches(13.333); presentation.slide_height = Inches(7.5)
                pages = []; heading = title; body = []
                for line in content.splitlines():
                    if re.match(r'^#{1,2}\s+', line):
                        if body: pages.append((heading, body)); body = []
                        heading = re.sub(r'^#+\s+', '', line)
                    elif line.strip():
                        text = line.strip().removeprefix('- ')
                        body += [text[i:i + 80] for i in range(0, len(text), 80)]
                if body or not pages: pages.append((heading, body or [title]))
                chunks = [(head, lines[i:i + 6]) for head, lines in pages for i in range(0, len(lines), 6)]
                if len(chunks) > 80: raise PolicyError('Presentation exceeds 80 slides')
                for number, (heading, lines) in enumerate(chunks, 1):
                    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor.from_string('F7F9F5')
                    box = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(11.9), Inches(1))
                    para = box.text_frame.paragraphs[0]; para.text = heading[:64]; para.font.size = Pt(28); para.font.bold = True; para.font.color.rgb = RGBColor.from_string('234E3A')
                    box = slide.shapes.add_textbox(Inches(.75), Inches(1.8), Inches(11.7), Inches(4.7)); box.text_frame.word_wrap = True
                    for i, line in enumerate(lines):
                        para = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                        para.text = line; para.font.size = Pt(19); para.space_after = Pt(14)
                    foot = slide.shapes.add_textbox(Inches(11.8), Inches(6.9), Inches(.7), Inches(.3)); foot.text_frame.text = str(number)
                presentation.save(target)
            with out.open('xb') as stream: stream.write(target.read_bytes())
        return {'path': relative, 'format': format, 'bytes': out.stat().st_size,
                'note': '新文件；基础文本与 Markdown 导出，不承诺复杂版式无损转换。'}


def register_documents(registry, workspace):
    from .tools import Tool, schema, TEXT
    tools = DocumentTools(workspace)
    registry.add(Tool('documents.read', 'Read text/tables from docx/xlsx/pptx/md. Paginated; no macro or formula execution.',
        schema({'path': TEXT, 'offset': {'type': 'integer', 'minimum': 0}, 'limit': {'type': 'integer', 'minimum': 1, 'maximum': 16000}}, ['path']), tools.read))
    registry.add(Tool('documents.create', 'Create a new md/docx/xlsx/pptx from text or Markdown. Never overwrite originals.',
        schema({'format': {'type': 'string', 'enum': list(FORMATS)}, 'content': {'type': 'string', 'minLength': 1, 'maxLength': MAX_TEXT}, 'title': {'type': 'string', 'minLength': 1, 'maxLength': 100}}, ['format', 'content']), tools.create))
